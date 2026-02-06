import asyncio
from pathlib import Path
import logging
from typing import Dict, Any, List
import PyPDF2
import docx

logger = logging.getLogger(__name__)

class DocumentParser:
    
    def __init__(self):
        self.supported_formats = {
            '.pdf': self._parse_pdf,
            '.docx': self._parse_docx,
            '.txt': self._parse_text
        }
    
    async def parse(self, file_path: Path) -> Dict[str, Any]:
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_extension = file_path.suffix.lower()
        
        if file_extension not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_extension}")
        
        parser_func = self.supported_formats[file_extension]
        
        try:
            loop = asyncio.get_event_loop()
            content = await loop.run_in_executor(None, parser_func, file_path)
        
            content['file_info'] = {
                'path': str(file_path),
                'size': file_path.stat().st_size,
                'modified': file_path.stat().st_mtime,
                'extension': file_extension
            }
            
            return content
            
        except Exception as e:
            logger.error(f"Error parsing {file_path}: {str(e)}")
            raise
    
    def _parse_pdf(self, file_path: Path) -> Dict[str, Any]:
  
        content = {
            'text': '',
            'metadata': {},
            'pages': [],
            'tables': []
        }
        
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            
            if pdf_reader.metadata:
                content['metadata'] = {
                    'title': pdf_reader.metadata.get('/Title', ''),
                    'author': pdf_reader.metadata.get('/Author', ''),
                    'subject': pdf_reader.metadata.get('/Subject', ''),
                    'creator': pdf_reader.metadata.get('/Creator', ''),
                    'producer': pdf_reader.metadata.get('/Producer', ''),
                    'creation_date': pdf_reader.metadata.get('/CreationDate', ''),
                    'modification_date': pdf_reader.metadata.get('/ModDate', '')
                }
            
            for page_num, page in enumerate(pdf_reader.pages, 1):
                page_text = page.extract_text()
                content['text'] += page_text + "\n\n"
                
                content['pages'].append({
                    'number': page_num,
                    'text': page_text,
                    'has_text': bool(page_text.strip())
                })
            
            content['file_info'] = {
                'total_pages': len(pdf_reader.pages),
                'is_encrypted': pdf_reader.is_encrypted
            }
        
        return content
    
    def _parse_docx(self, file_path: Path) -> Dict[str, Any]:
        content = {
            'text': '',
            'metadata': {},
            'pages': [],
            'tables': []
        }
        
        doc = docx.Document(file_path)
        
        core_props = doc.core_properties
        content['metadata'] = {
            'title': core_props.title or '',
            'author': core_props.author or '',
            'subject': core_props.subject or '',
            'keywords': core_props.keywords or '',
            'comments': core_props.comments or '',
            'created': str(core_props.created) if core_props.created else '',
            'modified': str(core_props.modified) if core_props.modified else '',
            'last_modified_by': core_props.last_modified_by or ''
        }
        
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        
        content['text'] = "\n".join(paragraphs)
  
        for table_num, table in enumerate(doc.tables, 1):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            
            if table_data:
                content['tables'].append({
                    'table_number': table_num,
                    'data': table_data,
                    'rows': len(table_data),
                    'columns': len(table_data[0]) if table_data else 0
                })
        
        content['pages'].append({
            'number': 1,
            'text': content['text'],
            'has_text': bool(content['text'].strip())
        })
        
        content['file_info'] = {
            'paragraphs': len(paragraphs),
            'tables': len(content['tables'])
        }
        
        return content
    
    def _parse_text(self, file_path: Path) -> Dict[str, Any]:

        content = {
            'text': '',
            'metadata': {},
            'pages': [],
            'tables': []
        }
        
        encodings = ['utf-8', 'latin-1', 'iso-8859-1', 'cp1252']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    text = file.read()
                    content['text'] = text
                    content['metadata']['encoding'] = encoding
                    break
            except UnicodeDecodeError:
                continue
        

        if not content['text']:
            with open(file_path, 'rb') as file:
                content['text'] = file.read().decode('utf-8', errors='ignore')
            content['metadata']['encoding'] = 'utf-8 (with errors ignored)'
        
        lines = content['text'].split('\n')
        lines_per_page = 50
        
        for i in range(0, len(lines), lines_per_page):
            page_lines = lines[i:i + lines_per_page]
            page_text = '\n'.join(page_lines)
            
            content['pages'].append({
                'number': (i // lines_per_page) + 1,
                'text': page_text,
                'has_text': bool(page_text.strip())
            })
        
        content['file_info'] = {
            'total_lines': len(lines),
            'total_pages': len(content['pages'])
        }
        
        return content
    
    def analyze_text(self, text: str) -> Dict[str, Any]:
        words = text.split()
        sentences = text.replace('!', '.').replace('?', '.').split('.')
        paragraphs = [p for p in text.split('\n\n') if p.strip()]
        
        return {
            'word_count': len(words),
            'sentence_count': len([s for s in sentences if s.strip()]),
            'paragraph_count': len(paragraphs),
            'avg_word_length': sum(len(word) for word in words) / len(words) if words else 0,
            'avg_sentence_length': len(words) / len(sentences) if sentences else 0
        }